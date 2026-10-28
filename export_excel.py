"""把核對後的單據匯出成 Excel（.xlsx）。

不依賴 Streamlit，可單獨測試。需要：pip install openpyxl
"""
import re
from datetime import datetime, date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SHEET_SUM = "單據彙總"
SHEET_ITEMS = "貨品明細"
SHEET_TODO = "未處理單據"
SHEET_INFO = "批次資訊"

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_EMOJI = re.compile(r"[\U00010000-\U0010ffff☀-➿️‍]")

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D0D5DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RED = PatternFill("solid", bgColor="FEE2E2", fgColor="FEE2E2")
YELLOW = PatternFill("solid", bgColor="FEF3C7", fgColor="FEF3C7")
GREEN = PatternFill("solid", bgColor="DCFCE7", fgColor="DCFCE7")
MONEY = "#,##0.00"


def clean_text(v) -> str:
    """去掉 Excel 不接受的控制字元與表情符號。"""
    return _EMOJI.sub("", _ILLEGAL.sub("", str(v))).strip()


def num(v, default=0.0) -> float:
    try:
        f = float(v)
        return default if f != f else f      # NaN -> default
    except (TypeError, ValueError):
        return default


def _set(ws, row, col, value, fmt=None):
    """寫入儲存格。文字一律當成純文字，避免 AI 輸出以 '=' 開頭被 Excel 當成公式執行。"""
    c = ws.cell(row=row, column=col)
    if isinstance(value, str):
        c.value = clean_text(value)
        c.data_type = "s"
    else:
        c.value = value
    if fmt:
        c.number_format = fmt
    c.border = BORDER
    c.alignment = Alignment(vertical="center", wrap_text=isinstance(value, str) and len(value) > 40)
    return c


def _formula(ws, row, col, f, fmt=None):
    c = ws.cell(row=row, column=col, value=f)
    if fmt:
        c.number_format = fmt
    c.border = BORDER
    c.alignment = Alignment(vertical="center")
    return c


def _header(ws, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _as_date(s):
    """YYYY-MM-DD 轉成真正的日期（Excel 可排序/篩選）；其他格式保留原文字。"""
    try:
        return datetime.strptime(str(s).strip(), "%Y-%m-%d").date()
    except ValueError:
        return clean_text(s)


def build_excel(docs, todo_rows, meta) -> bytes:
    """
    docs: list[dict]，每張要匯出的單據：
        source_file, doc_number, date, customer_name, stated_total,
        items: list[dict(item_name, quantity, unit_price, subtotal, is_unclear)],
        status, modified(bool), warnings: list[str], model, verified_by, timestamp
    todo_rows: list[dict(source_file, status, reason)]  失敗或未確認、沒有匯出的單據
    meta: dict，批次資訊（顯示在「批次資訊」工作表）
    """
    wb = Workbook()

    # ---------- 貨品明細（先建立，彙總要引用） ----------
    wi = wb.active
    wi.title = SHEET_ITEMS
    _header(wi, ["編號", "單號", "序號", "貨品／服務名稱", "數量", "單價", "小計（單據）",
                 "計算小計（數量×單價）", "小計差異", "手寫／模糊"],
            [8, 20, 7, 36, 10, 12, 14, 18, 12, 12])
    r = 2
    for no, d in enumerate(docs, 1):
        for seq, it in enumerate(d.get("items", []), 1):
            _set(wi, r, 1, no)
            _set(wi, r, 2, d.get("doc_number", ""))
            _set(wi, r, 3, seq)
            _set(wi, r, 4, it.get("item_name", ""))
            _set(wi, r, 5, num(it.get("quantity")), "#,##0.##")
            _set(wi, r, 6, num(it.get("unit_price")), MONEY)
            _set(wi, r, 7, num(it.get("subtotal")), MONEY)
            _formula(wi, r, 8, f"=ROUND(E{r}*F{r},2)", MONEY)
            _formula(wi, r, 9, f"=ROUND(G{r}-H{r},2)", MONEY)
            _set(wi, r, 10, "是" if it.get("is_unclear") else "否")
            r += 1
    last_i = max(r - 1, 2)
    wi.conditional_formatting.add(f"I2:I{last_i}", FormulaRule(formula=["ABS($I2)>0.01"], fill=RED))
    wi.conditional_formatting.add(f"J2:J{last_i}", CellIsRule(operator="equal", formula=['"是"'], fill=YELLOW))

    # ---------- 單據彙總 ----------
    ws = wb.create_sheet(SHEET_SUM, 0)
    _header(ws, ["編號", "來源檔案", "單號", "日期", "客戶", "品項數", "明細合計", "單據標示總金額",
                 "總額差異", "金額核對", "確認狀態", "人工修改過", "剩餘異常數", "異常說明",
                 "識別模型", "核對人員", "核對時間"],
            [8, 28, 20, 12, 28, 9, 14, 16, 12, 11, 22, 11, 11, 50, 24, 14, 20])
    for no, d in enumerate(docs, 1):
        r = no + 1
        warns = [clean_text(w) for w in d.get("warnings", [])]
        _set(ws, r, 1, no)
        _set(ws, r, 2, d.get("source_file", ""))
        _set(ws, r, 3, d.get("doc_number", ""))
        dv = _as_date(d.get("date", ""))
        _set(ws, r, 4, dv, "yyyy-mm-dd" if isinstance(dv, date) else None)
        _set(ws, r, 5, d.get("customer_name", ""))
        _formula(ws, r, 6, f"=COUNTIF('{SHEET_ITEMS}'!A:A,A{r})", "0")
        _formula(ws, r, 7, f"=SUMIF('{SHEET_ITEMS}'!A:A,A{r},'{SHEET_ITEMS}'!G:G)", MONEY)
        _set(ws, r, 8, num(d.get("stated_total")), MONEY)
        _formula(ws, r, 9, f"=ROUND(H{r}-G{r},2)", MONEY)
        _formula(ws, r, 10, f'=IF(ABS(I{r})<=0.01,"一致","不一致")')
        _set(ws, r, 11, d.get("status", ""))
        _set(ws, r, 12, "是" if d.get("modified") else "否")
        _set(ws, r, 13, len(warns))
        _set(ws, r, 14, "；".join(warns))
        _set(ws, r, 15, d.get("model", ""))
        _set(ws, r, 16, d.get("verified_by", ""))
        _set(ws, r, 17, d.get("timestamp", ""))
    last_s = max(len(docs) + 1, 2)
    ws.conditional_formatting.add(f"J2:J{last_s}", CellIsRule(operator="equal", formula=['"不一致"'], fill=RED))
    ws.conditional_formatting.add(f"J2:J{last_s}", CellIsRule(operator="equal", formula=['"一致"'], fill=GREEN))
    ws.conditional_formatting.add(f"L2:L{last_s}", CellIsRule(operator="equal", formula=['"是"'], fill=YELLOW))
    ws.conditional_formatting.add(f"M2:M{last_s}", CellIsRule(operator="greaterThan", formula=["0"], fill=RED))
    # 合計列
    tr = last_s + 2
    _set(ws, tr, 2, "合計")
    ws.cell(row=tr, column=2).font = Font(bold=True)
    _formula(ws, tr, 6, f"=SUM(F2:F{last_s})", "0")
    _formula(ws, tr, 7, f"=SUM(G2:G{last_s})", MONEY)
    _formula(ws, tr, 8, f"=SUM(H2:H{last_s})", MONEY)
    for col in (6, 7, 8):
        ws.cell(row=tr, column=col).font = Font(bold=True)

    # ---------- 未處理單據 ----------
    wt = wb.create_sheet(SHEET_TODO)
    _header(wt, ["來源檔案", "狀態", "原因／說明"], [32, 16, 90])
    for i, t in enumerate(todo_rows, 2):
        _set(wt, i, 1, t.get("source_file", ""))
        _set(wt, i, 2, t.get("status", ""))
        _set(wt, i, 3, t.get("reason", ""))
    if not todo_rows:
        _set(wt, 2, 1, "（沒有）")

    # ---------- 批次資訊 ----------
    wm = wb.create_sheet(SHEET_INFO)
    wm.column_dimensions["A"].width = 22
    wm.column_dimensions["B"].width = 60
    wm.cell(row=1, column=1, value="項目").fill = HEAD_FILL
    wm.cell(row=1, column=2, value="內容").fill = HEAD_FILL
    wm.cell(row=1, column=1).font = wm.cell(row=1, column=2).font = HEAD_FONT
    for i, (k, v) in enumerate(meta.items(), 2):
        _set(wm, i, 1, k)
        _set(wm, i, 2, v if isinstance(v, (int, float)) else str(v))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()